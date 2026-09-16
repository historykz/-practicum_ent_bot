"""
Персональная рабочая тетрадь: диагональная метка с ID ученика — как на
страницах конспекта. По утёкшему файлу видно, чья это копия.

PDF  — метка на каждой странице поверх содержимого (полупрозрачная, читать
       не мешает, вырезать нельзя — она часть страницы).
DOCX — метка в колонтитуле каждого раздела документа.
Остальные форматы отдаются как есть (в имени файла всё равно стоит ID).

Скачивание из Telegram Mini App идёт через встроенный Telegram.WebApp.
downloadFile: файл тянет сам клиент Telegram, без cookies сайта. Поэтому
ссылка подписывается коротким токеном на конкретного ученика и урок.
"""
import hashlib
import hmac
import io
import logging
import math
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Optional

log = logging.getLogger(__name__)

ALMATY = timezone(timedelta(hours=5))
TOKEN_TTL = 6 * 3600           # ссылка живёт 6 часов


def _secret() -> bytes:
    try:
        from webapp.server import _session_secret_bytes
        base = _session_secret_bytes()
    except Exception:
        import config
        base = hashlib.sha256(("smartent-session::" + config.BOT_TOKEN).encode()).digest()
    return hashlib.sha256(b"workbook::" + base).digest()


def make_token(lesson_id: int, tg_id: int, ttl: int = TOKEN_TTL) -> str:
    exp = int(time.time()) + int(ttl)
    msg = f"{int(lesson_id)}:{int(tg_id)}:{exp}".encode()
    sig = hmac.new(_secret(), msg, hashlib.sha256).hexdigest()[:24]
    return f"{int(tg_id)}.{exp}.{sig}"


def parse_token(lesson_id: int, token: str) -> Optional[int]:
    """tg_id, если подпись верна и срок не вышел; иначе None."""
    try:
        tg_s, exp_s, sig = (token or "").split(".")
        tg_id, exp = int(tg_s), int(exp_s)
    except (ValueError, AttributeError):
        return None
    if exp < time.time():
        return None
    msg = f"{int(lesson_id)}:{tg_id}:{exp}".encode()
    good = hmac.new(_secret(), msg, hashlib.sha256).hexdigest()[:24]
    return tg_id if hmac.compare_digest(good, sig) else None


def label(tg_id: int, username: Optional[str] = None) -> str:
    parts = []
    if username:
        parts.append(f"@{username}")
    parts.append(f"ID:{tg_id}")
    parts.append(datetime.now(ALMATY).strftime("%d.%m.%Y"))
    return "  -  ".join(parts)


def personal_name(filename: str, tg_id: int) -> str:
    p = Path(filename or "workbook.pdf")
    return f"{p.stem} (ID {tg_id}){p.suffix}"


# ───────────────────────── PDF ─────────────────────────

def _pdf_escape(text: str) -> str:
    safe = text.encode("latin-1", "replace").decode("latin-1")
    return safe.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _overlay_pdf(text: str, w: float, h: float) -> bytes:
    """Одностраничный PDF-слой: три диагональные надписи, серые, прозрачные.

    Собирается вручную (Helvetica — встроенный шрифт PDF), без reportlab.
    """
    angle = math.radians(28)
    c, s = round(math.cos(angle), 4), round(math.sin(angle), 4)
    size = max(12, round(min(w, h) / 24))
    esc = _pdf_escape(text)
    ops = ["/GS1 gs", "0.49 0.52 0.60 rg", f"BT /F1 {size} Tf"]
    for fx, fy in ((0.06, 0.16), (0.22, 0.50), (0.10, 0.82)):
        ops.append(f"{c} {s} {-s} {c} {round(w * fx, 1)} {round(h * fy, 1)} Tm ({esc}) Tj")
    ops.append("ET")
    content = "\n".join(ops).encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {w:.2f} {h:.2f}] "
         f"/Resources << /Font << /F1 4 0 R >> /ExtGState << /GS1 5 0 R >> >> "
         f"/Contents 6 0 R >>").encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        b"<< /Type /ExtGState /ca 0.38 /CA 0.38 >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def mark_pdf(data: bytes, text: str) -> bytes:
    from pypdf import PdfReader, PdfWriter
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("PDF защищён паролем")
    writer = PdfWriter()
    cache: dict = {}
    for page in reader.pages:
        box = page.mediabox
        w, h = float(box.width), float(box.height)
        key = (round(w), round(h))
        overlay = cache.get(key)
        if overlay is None:
            overlay = PdfReader(io.BytesIO(_overlay_pdf(text, w, h))).pages[0]
            cache[key] = overlay
        page.merge_page(overlay)
        writer.add_page(page)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


# ───────────────────────── DOCX ─────────────────────────

def mark_docx(data: bytes, text: str) -> bytes:
    import docx
    from docx.shared import Pt, RGBColor
    document = docx.Document(io.BytesIO(data))
    for section in document.sections:
        header = section.header
        header.is_linked_to_previous = False
        p = header.add_paragraph()
        run = p.add_run(text)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x7D, 0x85, 0x98)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# ───────────────────────── общий вход ─────────────────────────

def personalize(data: bytes, filename: str, tg_id: int,
                username: Optional[str] = None) -> tuple:
    """(байты, имя файла). При любой ошибке — исходный файл с ID в имени:
    ученик не должен остаться без тетради из-за экзотического PDF."""
    text = label(tg_id, username)
    ext = Path(filename or "").suffix.lower()
    name = personal_name(filename, tg_id)
    try:
        if ext == ".pdf":
            return mark_pdf(data, text), name
        if ext == ".docx":
            return mark_docx(data, text), name
    except Exception as e:
        log.warning("workbook mark %s: %s", filename, e)
    return data, name
